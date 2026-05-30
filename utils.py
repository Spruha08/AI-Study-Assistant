import os
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import chromadb
from sklearn.feature_extraction.text import HashingVectorizer
from langchain_groq import ChatGroq


def extract_text_from_pdf(file):
    """Extract text from a PDF file."""
    reader = PdfReader(file)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def extract_text_from_docx(file):
    """Extract text from a Word document."""
    document = Document(file)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    return text


def extract_text_from_pptx(file):
    """Extract text from a PowerPoint presentation."""
    presentation = Presentation(file)
    text = ""

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text += f"\nSlide {slide_number}:\n"

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"

    return text


def extract_text_from_file(file):
    """Detect file type and extract text."""
    file_name = file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_text_from_pdf(file)

    if file_name.endswith(".docx"):
        return extract_text_from_docx(file)

    if file_name.endswith(".pptx"):
        return extract_text_from_pptx(file)

    return ""


def split_text_into_chunks(text, chunk_size=1000, chunk_overlap=200):
    """Split long text into smaller overlapping chunks."""
    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]

        if chunk.strip():
            chunks.append(chunk)

        start += chunk_size - chunk_overlap

    return chunks


def get_vectorizer():
    """Create a lightweight text vectorizer that does not need Torch or ONNX."""
    return HashingVectorizer(
        n_features=384,
        alternate_sign=False,
        norm="l2",
    )


def create_vector_store(chunks, collection_name="study_materials"):
    """Create a ChromaDB knowledge base from text chunks."""
    client = chromadb.PersistentClient(path="chroma_db")

    try:
        client.delete_collection(name=collection_name)
    except Exception:
        pass

    collection = client.create_collection(name=collection_name)

    vectorizer = get_vectorizer()
    embeddings = vectorizer.transform(chunks).toarray().tolist()

    ids = [f"chunk_{i}" for i in range(len(chunks))]

    collection.add(
        documents=chunks,
        embeddings=embeddings,
        ids=ids,
    )

    return collection


def get_llm():
    """Load Groq Llama model using API key from .env."""
    load_dotenv()

    return ChatGroq(
        model="llama-3.3-70b-versatile",
        temperature=0.2,
        groq_api_key=os.getenv("GROQ_API_KEY"),
    )


def answer_question(collection, question, marks="5 marks"):
    """Answer a question using retrieved document chunks and Groq."""
    vectorizer = get_vectorizer()
    question_embedding = vectorizer.transform([question]).toarray().tolist()[0]

    results = collection.query(
        query_embeddings=[question_embedding],
        n_results=3,
    )

    retrieved_chunks = results["documents"][0]
    context = "\n\n".join(retrieved_chunks)

    llm = get_llm()

    prompt = f"""
You are an AI Study Assistant.

Answer the question using ONLY the provided study material context.
If the answer is not present in the context, say:
"I could not find this answer in the uploaded documents."

Marks requirement: {marks}

For 2 marks:
- Give a short and concise answer.

For 5 marks:
- Give a medium-length answer.
- Use clear bullet points if useful.

For 10 marks:
- Give a detailed structured answer with:
  - Introduction
  - Explanation
  - Examples if available
  - Advantages or disadvantages if relevant
  - Conclusion

Context:
{context}

Question:
{question}

Answer:
"""

    response = llm.invoke(prompt)
    return response.content


def generate_summary(collection, summary_type="Key Points"):
    """Generate a summary from uploaded study material."""
    results = collection.get()
    documents = results.get("documents", [])

    if not documents:
        return "No study material found. Please create the knowledge base first."

    context = "\n\n".join(documents[:8])

    llm = get_llm()

    prompt = f"""
You are an AI Study Assistant.

Generate a summary using ONLY the uploaded study material context.

Summary type: {summary_type}

Instructions:
- If summary type is "Chapter Summary", write a clear paragraph-wise summary.
- If summary type is "Quick Revision Notes", write short exam-focused notes.
- If summary type is "Key Points", write important points in bullets.
- Do not add information from outside the uploaded documents.
- Keep the language simple and student-friendly.

Context:
{context}

Summary:
"""

    response = llm.invoke(prompt)
    return response.content


def generate_quiz(collection, quiz_type="MCQs", number_of_questions=5):
    """Generate quiz questions from uploaded study material."""
    results = collection.get()
    documents = results.get("documents", [])

    if not documents:
        return "No study material found. Please create the knowledge base first."

    context = "\n\n".join(documents[:8])

    llm = get_llm()

    prompt = f"""
You are an AI Study Assistant.

Generate a quiz using ONLY the uploaded study material context.

Quiz type: {quiz_type}
Number of questions: {number_of_questions}

Instructions:
- If quiz type is "MCQs", create multiple choice questions with 4 options each and mention the correct answer.
- If quiz type is "Short Questions", create short answer questions with brief answers.
- If quiz type is "Viva Questions", create oral-exam style questions with expected answers.
- If quiz type is "Mixed Quiz", create a mix of MCQs, short questions, and viva questions with answers.
- Do not use information from outside the uploaded documents.
- Keep questions clear and exam-oriented.

Context:
{context}

Quiz:
"""

    response = llm.invoke(prompt)
    return response.content


def generate_interactive_mcqs(collection, number_of_questions=5):
    """Generate MCQs in a structured format from uploaded study material."""
    results = collection.get()
    documents = results.get("documents", [])

    if not documents:
        return []

    context = "\n\n".join(documents[:8])

    llm = get_llm()

    prompt = f"""
You are an AI Study Assistant.

Generate exactly {number_of_questions} multiple choice questions using ONLY the uploaded study material context.

Return the output in this exact format:

QUESTION: question text here
A: option A
B: option B
C: option C
D: option D
ANSWER: A

QUESTION: question text here
A: option A
B: option B
C: option C
D: option D
ANSWER: B

Rules:
- Use only the uploaded document context.
- Create clear exam-style MCQs.
- Each question must have exactly 4 options: A, B, C, D.
- The answer must be only one letter: A, B, C, or D.
- Do not add explanations.
- Do not add any extra heading or introduction.

Context:
{context}
"""

    response = llm.invoke(prompt)
    raw_text = response.content

    mcqs = []
    current_mcq = {}

    for line in raw_text.splitlines():
        line = line.strip()

        if line.startswith("QUESTION:"):
            if current_mcq:
                mcqs.append(current_mcq)

            current_mcq = {
                "question": line.replace("QUESTION:", "").strip(),
                "options": {},
                "answer": "",
            }

        elif line.startswith("A:"):
            current_mcq["options"]["A"] = line.replace("A:", "").strip()

        elif line.startswith("B:"):
            current_mcq["options"]["B"] = line.replace("B:", "").strip()

        elif line.startswith("C:"):
            current_mcq["options"]["C"] = line.replace("C:", "").strip()

        elif line.startswith("D:"):
            current_mcq["options"]["D"] = line.replace("D:", "").strip()

        elif line.startswith("ANSWER:"):
            current_mcq["answer"] = line.replace("ANSWER:", "").strip()

    if current_mcq:
        mcqs.append(current_mcq)

    valid_mcqs = []

    for mcq in mcqs:
        if (
            mcq.get("question")
            and len(mcq.get("options", {})) == 4
            and mcq.get("answer") in ["A", "B", "C", "D"]
        ):
            valid_mcqs.append(mcq)

    return valid_mcqs